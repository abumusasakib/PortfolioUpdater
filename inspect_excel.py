import openpyxl
from datetime import datetime

def inspect_excel(excel_path, output_file, sheets_file):
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    
    # Write sheet names to sheets.txt
    with open(sheets_file, 'w', encoding='utf-8') as f:
        for sheet_name in wb.sheetnames:
            f.write(f"{sheet_name}\n")
            
    # Write detailed inspection to excel_inspect.txt
    with open(output_file, 'w', encoding='utf-8') as f:
        f.write(f"Sheets: {wb.sheetnames}\n")
        
        target_sheets = ['01-personal', '02-career-highlight', '06-job-history', '12-contact', '07-project-roles']
        
        for sheet_name in target_sheets:
            if sheet_name in wb.sheetnames:
                f.write(f"\n--- Sheet: {sheet_name} ---\n")
                ws = wb[sheet_name]
                for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=20, values_only=True), 1):
                    f.write(f"Row {row_idx}: {row}\n")
            else:
                f.write(f"\nSheet {sheet_name} not found!\n")

if __name__ == "__main__":
    excel_path = "Résumé__Abu.Musa.Sakib.xlsx"
    inspect_excel(excel_path, "excel_inspect.txt", "sheets.txt")
