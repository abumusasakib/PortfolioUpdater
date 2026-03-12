import openpyxl
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_JUSTIFY
from reportlab.lib import colors
from datetime import datetime

# Function to extract data from Excel
def extract_excel_data(excel_path):
    """Extract resume data from Excel file"""
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    data = {}
    
    # Extract personal information
    ws = wb['01-personal']
    personal = {}
    for row in ws.iter_rows(min_row=3, values_only=True):
        if row[2] and row[3]:
            key = str(row[2]).strip()
            value = str(row[3]).strip() if row[3] else ""
            personal[key] = value
    data['personal'] = personal
    
    # Extract career highlights
    ws = wb['02-career-highlight']
    career_highlights = []
    for row in ws.iter_rows(min_row=3, values_only=True):
        # Career highlights are in column D (index 3) when column C (index 2) has a bullet
        if row[2] == '•' and row[3]:
            career_highlights.append(str(row[3]).strip())
    data['career_highlights'] = career_highlights
    
    # Extract education
    ws = wb['03-education']
    education = []
    for row in ws.iter_rows(min_row=4, values_only=True):
        if row[1] and row[2]:
            education.append({
                'year': str(row[1]).strip(),
                'degree': str(row[2]).strip(),
                'subject': str(row[3]).strip() if row[3] else "",
                'institute': str(row[4]).strip() if row[4] else ""
            })
    data['education'] = education
    
    # Extract technical expertise
    ws = wb['05-technical-expertise']
    technical_skills = []
    for row in ws.iter_rows(min_row=4, values_only=True):
        if row[1] and row[2] == '•':
            technical_skills.append({
                'category': str(row[1]).strip(),
                'skills': str(row[3]).strip() if row[3] else ""
            })
    data['technical_skills'] = technical_skills
    
    # Extract job history
    ws = wb['06-job-history']
    jobs = []
    current_job = None
    
    for row in ws.iter_rows(min_row=4, values_only=True):
        # Skip header row
        if row[1] == 'From':
            continue
            
        # Check if this is a date row (new job entry)
        if row[1] and isinstance(row[1], datetime):
            # Save previous job if exists
            if current_job and 'organization' in current_job and 'position' in current_job:
                jobs.append(current_job)
            
            # Start new job
            current_job = {
                'from': row[1].strftime('%Y-%m-%d') if isinstance(row[1], datetime) else str(row[1]),
                'to': row[2].strftime('%Y-%m-%d') if isinstance(row[2], datetime) else (str(row[2]) if row[2] and row[2] != 'Present' else 'Present'),
                'bullets': []
            }
            
        # Extract metadata and bullets (can be on the same row as date or separate)
        if current_job is not None and row[3]:
            text = str(row[3]).strip()
            
            if text == 'Organization':
                current_job['next_is_org'] = True
            elif text == 'Position':
                current_job['next_is_pos'] = True
            elif text == 'Job Summary':
                current_job['in_summary'] = True
            elif text == '•' and row[4]:
                # This is a bullet point
                current_job['bullets'].append(str(row[4]).strip())
            elif current_job.get('next_is_org'):
                current_job['organization'] = text
                current_job['next_is_org'] = False
            elif current_job.get('next_is_pos'):
                current_job['position'] = text
                current_job['next_is_pos'] = False
    
    # Don't forget the last job
    if current_job and 'organization' in current_job and 'position' in current_job:
        jobs.append(current_job)
    
    data['jobs'] = jobs
    
    # Extract contact info
    ws = wb['12-contact']
    contact = {}
    for row in ws.iter_rows(min_row=3, values_only=True):
        if row[1] and row[2]:
            key = str(row[1]).strip()
            value = str(row[2]).strip()
            contact[key] = value
    data['contact'] = contact
    
    # Extract projects
    # Sheet structure per project:
    #   Row: 'Project/Product'  → start of new project
    #   Row: '<project name>'   → name is in column D
    #   Row: 'Project Brief'    → next row is the summary
    #   Row: '<brief text>'     → description in column D
    #   Row: 'Tools and Technology' → following bullet rows have tech in col E
    #   Row: '•', tech          → tech value in col E
    #   Row: 'Activities/Tasks Performed' → following bullets are task descriptions
    #   Row: '•', task          → task text in col E
    projects = []
    if '07-project-roles' in wb.sheetnames:
        ws = wb['07-project-roles']
        current_project = None
        state = None  # track what the next row should contain

        CONTROL_LABELS = {
            'Project/Product', 'Client', 'Project Brief',
            'Tools and Technology', 'Role (s) Performed',
            'Activities/Tasks Performed'
        }

        for row in ws.iter_rows(min_row=4, values_only=True):
            d = str(row[3]).strip() if row[3] else ""
            e = str(row[4]).strip() if len(row) > 4 and row[4] else ""

            if d == 'Project/Product':
                if current_project and current_project.get('name'):
                    projects.append(current_project)
                current_project = {'name': '', 'description': '', 'tech': [], 'tasks': []}
                state = 'waiting_name'

            elif state == 'waiting_name' and d and d not in CONTROL_LABELS:
                current_project['name'] = d
                state = None

            elif d == 'Project Brief':
                state = 'brief'

            elif state == 'brief' and d and d not in CONTROL_LABELS:
                current_project['description'] = d
                state = None

            elif d == 'Tools and Technology':
                state = 'tech'

            elif state == 'tech' and d == '•' and e:
                # Split on commas that are NOT inside parentheses
                import re
                tech_items = [t.strip() for t in re.split(r',\s*(?![^(]*\))', e) if t.strip()]
                current_project['tech'].extend(tech_items)

            elif d == 'Activities/Tasks Performed':
                state = 'tasks'

            elif state == 'tasks' and d == '•' and e:
                current_project['tasks'].append(e)

            elif d in CONTROL_LABELS:
                state = None  # reset on any unrecognized control label

        if current_project and current_project.get('name'):
            projects.append(current_project)

    data['projects'] = projects
    
    return data