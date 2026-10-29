import argparse
import os
from datetime import datetime

import openpyxl
from bs4 import BeautifulSoup

from extract_excel_data import extract_excel_data


def public_location(address):
    """Keep a public portfolio at city/country precision, not street precision."""
    if not address:
        return 'Bangladesh'
    parts = [part.strip() for part in str(address).split(',') if part.strip()]
    if len(parts) >= 2:
        city = parts[-2].split('-')[0].strip()
        return f"{city}, {parts[-1]}"
    return parts[-1]


def format_phone(phone):
    digits = ''.join(char for char in str(phone or '') if char.isdigit())
    if digits.startswith('880') and len(digits) == 13:
        return f"+880 {digits[3:7]} {digits[7:]}"
    return f"+{digits}" if digits else ''


def append_text(soup, parent, tag_name, text, **attrs):
    tag = soup.new_tag(tag_name, attrs=attrs)
    tag.string = text
    parent.append(tag)
    return tag


def project_presentation(project_name):
    name = project_name.lower()
    if 'bank' in name:
        return ('bank', 'fa-building-columns')
    if 'ocr' in name:
        return ('ocr', 'fa-file-lines')
    if 'color' in name or 'fastapi' in name:
        return ('vision', 'fa-eye')
    return ('product', 'fa-gamepad')


def inspect_and_verify(excel_path):
    """Validate the expected resume sheets, then return the parsed portfolio data."""
    print('--- FOCUSED INSPECTION & VERIFICATION ---')
    print(f'Excel File: {excel_path}')

    workbook = openpyxl.load_workbook(excel_path, data_only=True)
    target_sheets = ['01-personal', '02-career-highlight', '06-job-history', '12-contact']
    print('\n[Row-Level Inspection Summary]')
    for sheet_name in target_sheets:
        if sheet_name in workbook.sheetnames:
            print(f"- Sheet '{sheet_name}': Found {workbook[sheet_name].max_row} rows.")
        else:
            print(f"- Sheet '{sheet_name}': NOT FOUND!")

    data = extract_excel_data(excel_path)
    print('\n[Data-Level Verification]')
    print(f"Name extracted: {data['personal'].get('Name')}")
    print(f"Designation: {data['personal'].get('Designation')}")
    print(f"Public email: {data['contact'].get('Email')}")
    print(f"Jobs found: {len(data['jobs'])}")
    print(f"Skills found: {len(data['technical_skills'])}")
    print(f"Projects found: {len(data['projects'])}")
    return data


def sync_common_elements(soup, data):
    """Synchronize only shared social/footer values across all pages."""
    linkedin = data['contact'].get('LinkedIn Page')
    github = data['contact'].get('GitHub/GitLab Page')
    for anchor in soup.select('.footer-socials a'):
        icon = anchor.find('i')
        icon_classes = ' '.join(icon.get('class', [])) if icon else ''
        if 'linkedin' in icon_classes and linkedin:
            anchor['href'] = linkedin
        elif 'github' in icon_classes and github:
            anchor['href'] = github

    footer_text = soup.select_one('.footer-text')
    if footer_text:
        name = data['personal'].get('Name', 'Abu Musa Sakib')
        footer_text.string = f'© {datetime.now().year} {name}'


def update_home(soup, data):
    personal = data['personal']
    name = personal.get('Name', 'Abu Musa Sakib')
    designation = personal.get('Designation', 'Software Developer')
    organization = personal.get('Organization', '')

    eyebrow = soup.select_one('.hero-intro .eyebrow')
    if eyebrow:
        eyebrow.string = f"{designation} / {public_location(personal.get('Address'))}"
    title = soup.select_one('#hero-title')
    if title:
        title.string = f'{name} builds thoughtful software.'
    copy = soup.select_one('.hero-copy')
    if copy:
        employer = f' at {organization}' if organization else ''
        copy.string = (
            f'I am a {designation.lower()}{employer}, focused on cross-platform Flutter '
            'applications, practical automation, and dependable delivery.'
        )

    side_note = soup.select_one('.hero-side-note')
    if side_note:
        skills = []
        for entry in data['technical_skills']:
            skills.extend(skill.strip() for skill in entry.get('skills', '').split(',') if skill.strip())
        side_note.clear()
        for index, skill in enumerate(skills[:4]):
            side_note.append(skill)
            if index < min(len(skills), 4) - 1:
                side_note.append(soup.new_tag('br'))


def update_about(soup, data):
    personal = data['personal']
    designation = personal.get('Designation', 'Software Engineer')
    organization = personal.get('Organization', '')
    summary = soup.select_one('.profile-summary')
    if summary:
        employer = f' at {organization}' if organization else ''
        summary.string = (
            f'I am a Bangladesh-based {designation.lower()}{employer} with a background in '
            'computer science and a focus on building useful, dependable software.'
        )

    details = [
        ('Based in', public_location(personal.get('Address'))),
        ('Nationality', personal.get('Nationality', 'Bangladeshi')),
    ]
    for row, (label, value) in zip(soup.select('.detail-list .detail-row'), details):
        term = row.find('dt')
        definition = row.find('dd')
        if term:
            term.string = label
        if definition:
            definition.string = value

    education_list = soup.select_one('.education-list')
    if not education_list:
        return
    heading = education_list.select_one('.eyebrow')
    education_list.clear()
    if heading:
        education_list.append(heading)
    for education in data['education']:
        article = soup.new_tag('article', attrs={'class': 'education-item'})
        append_text(soup, article, 'span', education.get('year', '').split('.')[0], **{'class': 'education-year'})
        body = soup.new_tag('div')
        degree = education.get('degree', '')
        subject = education.get('subject', '')
        append_text(soup, body, 'h3', f'{degree} in {subject}'.strip())
        append_text(soup, body, 'p', education.get('institute', ''))
        article.append(body)
        education_list.append(article)

    portrait = soup.select_one('.about-portrait')
    if portrait:
        portrait.attrs.pop('width', None)
        portrait.attrs.pop('height', None)


def update_skills(soup, data):
    skills_grid = soup.select_one('.skills-grid')
    if not skills_grid:
        return
    skills_grid.clear()
    for index, entry in enumerate(data['technical_skills'], 1):
        card = soup.new_tag('article', attrs={'class': 'skill-card'})
        append_text(soup, card, 'p', f'{index:02d} / Expertise', **{'class': 'skill-kicker'})
        append_text(soup, card, 'h2', entry.get('category', 'Skill'), **{'class': 'skill-title'})
        skill_list = soup.new_tag('ul', attrs={'class': 'skill-list'})
        for skill in entry.get('skills', '').split(','):
            if skill.strip():
                append_text(soup, skill_list, 'li', skill.strip())
        card.append(skill_list)
        skills_grid.append(card)


def update_projects(soup, data):
    projects_grid = soup.select_one('.projects-grid')
    if not projects_grid or not data.get('projects'):
        return
    projects_grid.clear()
    for index, project in enumerate(data['projects'], 1):
        presentation, icon = project_presentation(project['name'])
        card = soup.new_tag('article', attrs={'class': f'project-card project-card--{presentation}'})
        visual = soup.new_tag('div', attrs={'class': 'project-visual', 'aria-hidden': 'true'})
        append_text(soup, visual, 'span', f'{index:02d}', **{'class': 'project-index'})
        visual.append(soup.new_tag('i', attrs={'class': f'fa-solid {icon} project-icon'}))
        content = soup.new_tag('div', attrs={'class': 'project-content'})
        meta = soup.new_tag('div', attrs={'class': 'project-meta'})
        append_text(soup, meta, 'p', project.get('role') or 'Project experience', **{'class': 'mono-label'})
        period = ' - '.join(value for value in (project.get('from'), project.get('to')) if value)
        append_text(soup, meta, 'p', period, **{'class': 'project-period'})
        content.append(meta)
        append_text(soup, content, 'h2', project['name'], **{'class': 'project-title'})
        append_text(soup, content, 'p', project.get('description', '').strip(), **{'class': 'project-description'})
        facts = soup.new_tag('dl', attrs={'class': 'project-facts'})
        if project.get('client'):
            client = soup.new_tag('div')
            append_text(soup, client, 'dt', 'Client')
            append_text(soup, client, 'dd', project['client'])
            facts.append(client)
        if project.get('tasks'):
            focus = soup.new_tag('div')
            append_text(soup, focus, 'dt', 'Focus')
            append_text(soup, focus, 'dd', project['tasks'][0])
            facts.append(focus)
        content.append(facts)
        technologies = soup.new_tag('div', attrs={'class': 'project-tech'})
        for technology in project.get('tech', []):
            append_text(soup, technologies, 'span', technology, **{'class': 'tech-badge'})
        content.append(technologies)
        card.append(visual)
        card.append(content)
        projects_grid.append(card)


def update_contact(soup, data):
    email = data['contact'].get('Email', '')
    phone = format_phone(data['contact'].get('Contact Number', ''))
    email_anchor = soup.select_one('a[href^="mailto:"]')
    if email_anchor and email:
        email_anchor['href'] = f'mailto:{email}'
        email_anchor.string = email
    phone_anchor = soup.select_one('a[href^="tel:"]')
    if phone_anchor and phone:
        phone_anchor['href'] = f"tel:{phone.replace(' ', '')}"
        phone_anchor.string = phone
    for anchor in soup.select('.contact-detail a'):
        label = anchor.get_text(strip=True).lower()
        if label == 'linkedin':
            anchor['href'] = data['contact'].get('LinkedIn Page', '')
        elif label == 'github':
            anchor['href'] = data['contact'].get('GitHub/GitLab Page', '')


def update_website(data, base_dir):
    if not os.path.isdir(base_dir):
        raise FileNotFoundError(f'Portfolio directory not found: {base_dir}')

    updates = {
        'index.html': update_home,
        'about-me.html': update_about,
        'skills.html': update_skills,
        'portfolio.html': update_projects,
        'contact.html': update_contact,
    }
    print(f'\n--- UPDATING WEBSITE IN {base_dir} ---')
    for page_name, page_update in updates.items():
        page_path = os.path.join(base_dir, page_name)
        if not os.path.isfile(page_path):
            print(f'Skipping missing page: {page_name}')
            continue
        print(f'Processing {page_name}...')
        with open(page_path, encoding='utf-8') as page_file:
            soup = BeautifulSoup(page_file, 'html.parser')
        sync_common_elements(soup, data)
        page_update(soup, data)
        with open(page_path, 'w', encoding='utf-8') as page_file:
            page_file.write(str(soup))
    print('\nWebsite updated successfully!')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='Synchronize the portfolio from a resume workbook.')
    parser.add_argument('--excel', default='Résumé__Abu.Musa.Sakib.xlsx', help='Path to the resume workbook.')
    parser.add_argument('--portfolio', default='portfolio-sakib', help='Path to the portfolio repository.')
    args = parser.parse_args()
    resume_data = inspect_and_verify(args.excel)
    update_website(resume_data, args.portfolio)
